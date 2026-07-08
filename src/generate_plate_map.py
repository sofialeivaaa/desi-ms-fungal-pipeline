"""
Genera un libro de Excel (.xlsx) para el seguimiento de una microplaca de 384 pocillos
(Eppendorf Twin.tec 384) usada en el pipeline DESI-MS de hongos.

Lee las muestras reales desde:
    data/plate_inputs/{PLATE_ID}_input.xlsx   (pestaña "input")

Genera:
    data/plates/{PLATE_ID}.xlsx
        - Plate_Map        -> cuadrícula visual de la placa (16 filas A-P x 24 columnas)
        - Sample_Metadata  -> tabla plana (384 filas) lista para pandas

Requiere: openpyxl

Uso:
    python src/generate_plate_map.py PLATE_01
    python src/generate_plate_map.py PLATE_02
"""
import sys
import string
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

FONT_NAME = "Calibri"
PLATE_ID = sys.argv[1] if len(sys.argv) > 1 else "PLATE_01"

REPO_ROOT = Path(__file__).resolve().parent.parent
INPUT_PATH = REPO_ROOT / "data" / "plate_inputs" / f"{PLATE_ID}_input.xlsx"
OUTPUT_DIR = REPO_ROOT / "data" / "plates"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

ROWS = list(string.ascii_uppercase[:16])   # A ... P
COLS = list(range(1, 25))                  # 1 ... 24
VALID_WELLS = {f"{r}{c}" for r in ROWS for c in COLS}

# ---------------------------------------------------------------------------
# 1) Leer la plantilla de entrada llenada a mano
# ---------------------------------------------------------------------------
if not INPUT_PATH.exists():
    print(f"ERROR: no encontré el archivo de entrada:\n  {INPUT_PATH}")
    print("Crea ese archivo (copia PLATE_input_template.xlsx, llénalo, y guárdalo con ese nombre exacto).")
    sys.exit(1)

wb_in = load_workbook(INPUT_PATH, data_only=True)
if "input" not in wb_in.sheetnames:
    print(f"ERROR: el archivo {INPUT_PATH.name} no tiene una pestaña llamada 'input'.")
    sys.exit(1)

ws_in = wb_in["input"]
header = [c.value for c in ws_in[1]]
required_cols = ["well", "sample_id", "organism", "condition", "replicate", "notes"]
missing = [col for col in required_cols if col not in header]
if missing:
    print(f"ERROR: faltan columnas en la plantilla de entrada: {missing}")
    sys.exit(1)
col_idx = {name: header.index(name) for name in required_cols}

EXAMPLES = {}
errors = []
for row in ws_in.iter_rows(min_row=2, values_only=True):
    if row[col_idx["well"]] is None:
        continue
    well = str(row[col_idx["well"]]).strip().upper()
    if well == "":
        continue
    if well not in VALID_WELLS:
        errors.append(f"Pocillo inválido: '{well}' (debe ser letra A-P + número 1-24)")
        continue
    if well in EXAMPLES:
        errors.append(f"Pocillo duplicado en la plantilla: '{well}'")
        continue
    sample_id = row[col_idx["sample_id"]] or ""
    organism = row[col_idx["organism"]] or ""
    condition = row[col_idx["condition"]] or ""
    replicate = row[col_idx["replicate"]] if row[col_idx["replicate"]] is not None else ""
    notes = row[col_idx["notes"]] or ""
    EXAMPLES[well] = (sample_id, organism, condition, replicate, notes)

if errors:
    print("ERROR: se encontraron problemas en la plantilla de entrada:")
    for e in errors:
        print(f"  - {e}")
    sys.exit(1)

if not EXAMPLES:
    print(f"ERROR: {INPUT_PATH.name} no tiene ninguna fila de muestra llenada.")
    sys.exit(1)

# Etiqueta corta para el mapa visual: usa el sample_id (o el well si no hay dato)
PLATE_LABEL = {well: str(data[0]) if data[0] else well for well, data in EXAMPLES.items()}

# ---------------------------------------------------------------------------
# 2) Construir el libro de salida
# ---------------------------------------------------------------------------
wb = Workbook()

# ===== PLATE_MAP =====
ws_map = wb.active
ws_map.title = "Plate_Map"

title_font   = Font(name=FONT_NAME, size=14, bold=True, color="1F1F1F")
axis_font    = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
well_font    = Font(name=FONT_NAME, size=8, color="1F1F1F")
example_font = Font(name=FONT_NAME, size=8, bold=True, color="1F1F1F")

axis_fill    = PatternFill("solid", fgColor="404040")
empty_fill   = PatternFill("solid", fgColor="FFFFFF")
sample_fill  = PatternFill("solid", fgColor="BDD7EE")
control_fill = PatternFill("solid", fgColor="F8CBAD")

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws_map.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS) + 1)
ws_map.cell(row=1, column=1, value=f"{PLATE_ID} — Mapa de Placa (384 pocillos, 16x24)").font = title_font
ws_map.row_dimensions[1].height = 22

header_row = 3
first_data_row = 4

ws_map.cell(row=header_row, column=1, value="").fill = axis_fill
ws_map.cell(row=header_row, column=1).border = border
for j, col_num in enumerate(COLS, start=2):
    c = ws_map.cell(row=header_row, column=j, value=col_num)
    c.font = axis_font
    c.fill = axis_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border

for i, row_letter in enumerate(ROWS, start=0):
    r = first_data_row + i
    rh = ws_map.cell(row=r, column=1, value=row_letter)
    rh.font = axis_font
    rh.fill = axis_fill
    rh.alignment = Alignment(horizontal="center", vertical="center")
    rh.border = border

    for j, col_num in enumerate(COLS, start=2):
        well = f"{row_letter}{col_num}"
        c = ws_map.cell(row=r, column=j)
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.font = example_font if well in PLATE_LABEL else well_font
        if well in PLATE_LABEL:
            c.value = PLATE_LABEL[well]
            condition_val = str(EXAMPLES[well][2]).lower()
            c.fill = control_fill if "control" in condition_val else sample_fill
        else:
            c.fill = empty_fill

ws_map.column_dimensions["A"].width = 4
for j in range(2, len(COLS) + 2):
    ws_map.column_dimensions[get_column_letter(j)].width = 11
for i in range(len(ROWS)):
    ws_map.row_dimensions[first_data_row + i].height = 26

ws_map.freeze_panes = ws_map.cell(row=first_data_row, column=2)

# ===== SAMPLE_METADATA =====
ws_meta = wb.create_sheet("Sample_Metadata")

columns = ["plate_id", "well", "sample_id", "organism", "condition", "replicate", "filename", "notes"]
widths  = [12, 8, 12, 24, 24, 10, 26, 30]

header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="404040")
cell_font   = Font(name=FONT_NAME, size=11, color="1F1F1F")

for i, h in enumerate(columns, start=1):
    c = ws_meta.cell(row=1, column=i, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border
    ws_meta.column_dimensions[get_column_letter(i)].width = widths[i-1]

r = 2
filled_count = 0
for row_letter in ROWS:
    for col_num in COLS:
        well = f"{row_letter}{col_num}"
        filename = f"{PLATE_ID}_{well}.mzML"
        if well in EXAMPLES:
            sample_id, organism, condition, replicate, notes = EXAMPLES[well]
            filled_count += 1
        else:
            sample_id, organism, condition, replicate, notes = ("", "", "", "", "")

        rowdata = [PLATE_ID, well, sample_id, organism, condition, replicate, filename, notes]
        for i, val in enumerate(rowdata, start=1):
            c = ws_meta.cell(row=r, column=i, value=val)
            c.font = cell_font
            c.border = border
            c.alignment = Alignment(horizontal="left" if i not in (2, 6) else "center", vertical="center")
        r += 1

last_row = r - 1
last_col_letter = get_column_letter(len(columns))
table_ref = f"A1:{last_col_letter}{last_row}"
tab = Table(displayName="SampleMetadata384", ref=table_ref)
tab.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2", showFirstColumn=False,
    showLastColumn=False, showRowStripes=True, showColumnStripes=False
)
ws_meta.add_table(tab)
ws_meta.freeze_panes = "A2"

output_path = OUTPUT_DIR / f"{PLATE_ID}.xlsx"
wb.save(output_path)
print(f"Listo: {filled_count} pocillos llenados de 384 totales.")
print(f"Guardado en: {output_path}")